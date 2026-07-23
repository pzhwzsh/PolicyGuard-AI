"""Auditable CSV/XLSX field mapping and deterministic cleaning before batch review."""

from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

CANONICAL_FIELDS = ("external_id", "category", "title", "description", "markets")
REQUIRED_FIELDS = {"title", "markets"}
MAX_TABLE_ROWS = 1000
MAX_TABLE_COLUMNS = 100
HEADER_SCAN_ROWS = 20
FIELD_ALIASES = {
    "external_id": {"externalid", "sku", "skuid", "商品id", "商品编号", "产品编号", "货号"},
    "category": {"category", "类目", "商品类目", "产品类目", "分类"},
    "title": {"title", "商品标题", "商品名称", "产品名称", "品名", "标题"},
    "description": {"description", "商品描述", "产品描述", "描述", "卖点", "详情"},
    "markets": {"markets", "market", "市场", "目标市场", "销售市场", "国家", "站点"},
}
MARKET_ALIASES = {
    "CN": "CN", "中国": "CN", "中国大陆": "CN", "CHINA": "CN",
    "US": "US", "USA": "US", "美国": "US", "UNITEDSTATES": "US",
    "EU": "EU", "欧盟": "EU", "欧洲联盟": "EU", "EUROPEANUNION": "EU",
}


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _header_key(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", str(value or "").casefold())


def _canonical_header(value: Any) -> str | None:
    key = _header_key(value)
    return next((field for field, aliases in FIELD_ALIASES.items() if key in aliases), None)


def _clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _safe_cell(value: Any) -> Any:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else value


def _market_values(value: Any) -> tuple[list[str], list[str]]:
    raw = _clean_text(value)
    parts = [
        item.strip() for item in re.split(r"[,，;；/|、]+", raw)
        if item.strip()
    ]
    markets: list[str] = []
    invalid: list[str] = []
    for item in parts:
        key = re.sub(r"\s+", "", item).upper()
        market = MARKET_ALIASES.get(key)
        if market is None:
            invalid.append(item)
        elif market not in markets:
            markets.append(market)
    return markets, invalid


def _bounded_row(row: list[Any]) -> list[Any]:
    if len(row) > MAX_TABLE_COLUMNS:
        raise ValueError("table_column_limit_exceeded")
    return row


def _read_csv(path: Path, max_rows: int) -> list[tuple[str, list[list[Any]]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = []
        for index, row in enumerate(csv.reader(handle)):
            if index >= max_rows + HEADER_SCAN_ROWS + 1:
                raise ValueError("table_row_limit_exceeded")
            rows.append(_bounded_row(list(row)))
        return [("CSV", rows)]


def _read_xlsx(path: Path, max_rows: int) -> list[tuple[str, list[list[Any]]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            if sheet.max_column > MAX_TABLE_COLUMNS:
                raise ValueError("table_column_limit_exceeded")
            if sheet.max_row > max_rows + HEADER_SCAN_ROWS:
                raise ValueError("table_row_limit_exceeded")
            rows = [_bounded_row(list(row)) for row in sheet.iter_rows(values_only=True)]
            sheets.append((sheet.title, rows))
        return sheets
    finally:
        workbook.close()


def _header_index(rows: list[list[Any]]) -> int:
    candidates = []
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        values = [value for value in row if _clean_text(value)]
        mapped = sum(_canonical_header(value) is not None for value in values)
        candidates.append((mapped, len(set(map(_header_key, values))), -index, index))
    if not candidates or max(candidates)[0] == 0:
        raise ValueError("table_header_not_detected")
    return max(candidates)[-1]


def _sheet_mapping(
    headers: list[str], overrides: dict[str, str] | None = None,
) -> tuple[dict[str, str], list[dict[str, Any]]]:
    override_lookup = {_header_key(key): value for key, value in (overrides or {}).items()}
    mapping: dict[str, str] = {}
    issues: list[dict[str, Any]] = []
    claimed: set[str] = set()
    for header in headers:
        canonical = override_lookup.get(_header_key(header), _canonical_header(header))
        if canonical in {None, "ignore", ""}:
            continue
        if canonical not in CANONICAL_FIELDS:
            issues.append({"code": "unknown_mapping_target", "severity": "error", "column": header})
            continue
        if canonical in claimed:
            issues.append({
                "code": "duplicate_mapping_target", "severity": "error",
                "column": header, "field": canonical,
            })
            continue
        mapping[header] = canonical
        claimed.add(canonical)
    for field in sorted(REQUIRED_FIELDS - claimed):
        issues.append({"code": "required_field_unmapped", "severity": "error", "field": field})
    return mapping, issues


def analyze_table(
    path: Path,
    *,
    mapping_overrides: dict[str, dict[str, str]] | None = None,
    max_rows: int = 1000,
) -> dict[str, Any]:
    sheets = (
        _read_csv(path, max_rows)
        if path.suffix.lower() == ".csv"
        else _read_xlsx(path, max_rows)
    )
    analyzed_sheets = []
    skipped_sheets = []
    cleaned_rows = []
    all_issues = []
    seen_ids: set[str] = set()
    total_rows = 0
    for sheet_name, matrix in sheets:
        if not matrix or not any(any(_clean_text(cell) for cell in row) for row in matrix):
            continue
        try:
            header_row = _header_index(matrix)
        except ValueError as exc:
            if str(exc) != "table_header_not_detected":
                raise
            skipped_sheets.append({"name": sheet_name, "reason": str(exc)})
            continue
        width = max(len(row) for row in matrix)
        raw_headers = [
            _clean_text(matrix[header_row][index] if index < len(matrix[header_row]) else "")
            or f"column_{index + 1}"
            for index in range(width)
        ]
        headers = []
        header_counts: dict[str, int] = {}
        for header in raw_headers:
            count = header_counts.get(header, 0) + 1
            header_counts[header] = count
            headers.append(header if count == 1 else f"{header}__{count}")
        mapping, mapping_issues = _sheet_mapping(
            headers, (mapping_overrides or {}).get(sheet_name)
        )
        sheet_issues = [dict(item, sheet=sheet_name) for item in mapping_issues]
        sheet_rows = []
        for source_row, values in enumerate(matrix[header_row + 1:], start=header_row + 2):
            if not any(_clean_text(value) for value in values):
                continue
            total_rows += 1
            if total_rows > max_rows:
                raise ValueError("table_row_limit_exceeded")
            raw = {
                header: _json_value(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
            }
            canonical_raw = {canonical: raw.get(header) for header, canonical in mapping.items()}
            issues: list[dict[str, Any]] = []
            for header, value in raw.items():
                if isinstance(value, str) and value.startswith("="):
                    issues.append({
                        "code": "formula_cell_not_evaluated", "severity": "warning",
                        "column": header,
                    })
            external_id = _clean_text(canonical_raw.get("external_id"))
            if not external_id:
                external_id = f"{_header_key(sheet_name).upper() or 'SHEET'}-{source_row}"
                issues.append({"code": "external_id_generated", "severity": "warning"})
            if external_id in seen_ids:
                issues.append({"code": "duplicate_external_id", "severity": "error"})
            seen_ids.add(external_id)
            title = _clean_text(canonical_raw.get("title"))
            if not title:
                issues.append({"code": "title_required", "severity": "error", "field": "title"})
            category = _clean_text(canonical_raw.get("category")) or "all"
            description = _clean_text(canonical_raw.get("description"))
            markets, invalid_markets = _market_values(canonical_raw.get("markets"))
            if invalid_markets:
                issues.append({
                    "code": "invalid_markets", "severity": "error",
                    "field": "markets", "values": invalid_markets,
                })
            if not markets:
                issues.append({"code": "markets_required", "severity": "error", "field": "markets"})
            cleaned = {
                "external_id": external_id,
                "category": category,
                "title": title,
                "description": description,
                "markets": ",".join(markets),
            }
            item = {
                "source_sheet": sheet_name,
                "source_row": source_row,
                "raw": raw,
                "cleaned": cleaned,
                "issues": issues,
                "valid": not any(issue["severity"] == "error" for issue in issues),
            }
            sheet_rows.append(item)
            cleaned_rows.append(item)
            all_issues.extend(dict(issue, sheet=sheet_name, row=source_row) for issue in issues)
        analyzed_sheets.append({
            "name": sheet_name,
            "header_row": header_row + 1,
            "columns": headers,
            "mapping": mapping,
            "issues": sheet_issues,
            "row_count": len(sheet_rows),
        })
        all_issues.extend(sheet_issues)
    if not analyzed_sheets:
        raise ValueError("table_has_no_data")
    error_count = sum(issue["severity"] == "error" for issue in all_issues)
    warning_count = sum(issue["severity"] == "warning" for issue in all_issues)
    mapping_error_count = sum(
        issue["severity"] == "error"
        for sheet in analyzed_sheets
        for issue in sheet["issues"]
    )
    return {
        "sheets": analyzed_sheets,
        "skipped_sheets": skipped_sheets,
        "rows": cleaned_rows,
        "issues": all_issues,
        "summary": {
            "sheet_count": len(analyzed_sheets),
            "skipped_sheet_count": len(skipped_sheets),
            "row_count": len(cleaned_rows),
            "valid_row_count": sum(item["valid"] for item in cleaned_rows),
            "invalid_row_count": sum(not item["valid"] for item in cleaned_rows),
            "error_count": error_count,
            "mapping_error_count": mapping_error_count,
            "warning_count": warning_count,
            "deterministic_row_count": len(cleaned_rows),
            "model_call_count": 0,
            "estimated_model_tokens": 0,
        },
    }


class TableCleaningWorkspace:
    def __init__(self, root: Path) -> None:
        self.root = root

    def stage(self, content: bytes, filename: str) -> dict[str, Any]:
        suffix = Path(filename).suffix.lower()
        if suffix not in {".csv", ".xlsx"}:
            raise ValueError("batch_file_type_not_supported")
        table_id = sha256(content).hexdigest()[:20]
        directory = self.root / table_id
        directory.mkdir(parents=True, exist_ok=True)
        original = directory / f"original{suffix}"
        if not original.exists():
            original.write_bytes(content)
        try:
            analysis = analyze_table(original)
        except ValueError:
            raise
        except (
            BadZipFile, InvalidFileException, UnicodeError, csv.Error,
            EOFError, KeyError, OSError,
        ) as exc:
            raise ValueError("batch_file_invalid") from exc
        manifest_path = directory / "manifest.json"
        existing = (
            json.loads(manifest_path.read_text(encoding="utf-8"))
            if manifest_path.exists()
            else {}
        )
        manifest = {
            "table_id": table_id,
            "filename": filename,
            "original_path": str(original.resolve()),
            "status": "preview",
            "revision": int(existing.get("revision", 0)),
            "content_hash": sha256(content).hexdigest(),
        }
        (directory / "analysis.json").write_text(
            json.dumps(analysis, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return {"table_id": table_id, "manifest": manifest, "analysis": analysis}

    def load(self, table_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
        if not re.fullmatch(r"[a-f0-9]{20}", table_id):
            raise LookupError("table_cleaning_not_found")
        directory = self.root / table_id
        manifest_path = directory / "manifest.json"
        analysis_path = directory / "analysis.json"
        if not manifest_path.is_file() or not analysis_path.is_file():
            raise LookupError("table_cleaning_not_found")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("table_id") != table_id:
            raise LookupError("table_cleaning_not_found")
        return manifest, json.loads(analysis_path.read_text(encoding="utf-8"))

    def _original_path(self, table_id: str, manifest: dict[str, Any]) -> Path:
        directory = (self.root / table_id).resolve()
        original = Path(str(manifest.get("original_path", ""))).resolve()
        if directory not in original.parents or not original.is_file():
            raise RuntimeError("table_cleaning_source_invalid")
        if original.name not in {"original.csv", "original.xlsx"}:
            raise RuntimeError("table_cleaning_source_invalid")
        return original

    def confirm(
        self,
        table_id: str,
        *,
        expected_revision: int,
        reviewer: str,
        field_mappings: dict[str, dict[str, str]],
        allow_partial: bool,
    ) -> dict[str, Any]:
        manifest, _ = self.load(table_id)
        if int(manifest.get("revision", 0)) != expected_revision:
            raise RuntimeError("table_cleaning_revision_conflict")
        if manifest.get("status") != "preview":
            raise RuntimeError("table_cleaning_already_confirmed")
        original = self._original_path(table_id, manifest)
        analysis = analyze_table(original, mapping_overrides=field_mappings)
        if analysis["summary"]["mapping_error_count"]:
            raise RuntimeError("table_cleaning_mapping_invalid")
        if analysis["summary"]["invalid_row_count"] and not allow_partial:
            raise RuntimeError("table_cleaning_has_blocking_errors")
        directory = self.root / table_id
        valid_rows = [item for item in analysis["rows"] if item["valid"]]
        if not valid_rows:
            raise RuntimeError("table_cleaning_has_no_valid_rows")
        cleaned_path = directory / "cleaned-input.csv"
        with cleaned_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(CANONICAL_FIELDS))
            writer.writeheader()
            writer.writerows(
                {
                    key: _safe_cell(value)
                    for key, value in item["cleaned"].items()
                }
                for item in valid_rows
            )
        report_path = directory / "cleaning-report.csv"
        with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=[
                    "source_sheet", "source_row", "valid", "issue_codes",
                    *CANONICAL_FIELDS,
                ],
            )
            writer.writeheader()
            for item in analysis["rows"]:
                writer.writerow({
                    "source_sheet": _safe_cell(item["source_sheet"]),
                    "source_row": item["source_row"],
                    "valid": item["valid"],
                    "issue_codes": ",".join(issue["code"] for issue in item["issues"]),
                    **{key: _safe_cell(value) for key, value in item["cleaned"].items()},
                })
        revision = expected_revision + 1
        manifest.update({
            "status": "confirmed",
            "revision": revision,
            "reviewer": reviewer,
            "allow_partial": allow_partial,
            "cleaned_path": str(cleaned_path.resolve()),
            "report_path": str(report_path.resolve()),
            "confirmed_mappings": field_mappings,
        })
        (directory / "analysis.json").write_text(
            json.dumps(analysis, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (directory / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return {"table_id": table_id, "manifest": manifest, "analysis": analysis}
