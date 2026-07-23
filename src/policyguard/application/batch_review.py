import csv
import json
from pathlib import Path

from sqlalchemy.orm import Session

from policyguard.application.workflow import ComplianceWorkflowService
from policyguard.infrastructure.repositories import (
    SqlAlchemyKnowledgeRepository,
    SqlAlchemyWorkflowRepository,
)

REQUIRED_COLUMNS = {"external_id", "category", "title", "description", "markets"}


def _normalize_row(row: dict, row_number: int) -> dict:
    normalized = {
        str(key or "").strip(): "" if value is None else str(value).strip()
        for key, value in row.items()
    }
    missing = REQUIRED_COLUMNS - normalized.keys()
    if missing:
        raise ValueError(f"batch_missing_columns:{','.join(sorted(missing))}")
    if not normalized["title"]:
        raise ValueError(f"batch_title_required:{row_number}")
    market_text = normalized["markets"].replace("，", ",").replace(";", ",").replace("；", ",")
    markets = [item.strip().upper() for item in market_text.split(",")]
    markets = [item for item in markets if item]
    if not markets or any(item not in {"CN", "US", "EU"} for item in markets):
        raise ValueError(f"batch_markets_invalid:{row_number}")
    return {
        "row_number": row_number,
        "external_id": normalized["external_id"] or f"ROW-{row_number}",
        "category": normalized["category"] or "all",
        "title": normalized["title"],
        "description": normalized["description"],
        "markets": markets,
    }


def read_batch_rows(path: Path, *, max_rows: int = 1000) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            raw_rows = list(csv.DictReader(handle))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [str(item or "").strip() for item in next(values, ())]
            raw_rows = [
                dict(zip(headers, row, strict=False))
                for row in values
                if any(item is not None and str(item).strip() for item in row)
            ]
        finally:
            workbook.close()
    else:
        raise ValueError("batch_file_type_not_supported")
    if not raw_rows:
        raise ValueError("batch_file_empty")
    if len(raw_rows) > max_rows:
        raise ValueError("batch_row_limit_exceeded")
    return [_normalize_row(row, index + 2) for index, row in enumerate(raw_rows)]


def _safe_cell(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def process_batch_review(
    session: Session,
    input_path: Path,
    output_dir: Path,
) -> dict:
    rows = read_batch_rows(input_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = output_dir / "results.jsonl"
    results = []
    completed_rows = set()
    if checkpoint_path.exists():
        for line in checkpoint_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                item = json.loads(line)
                results.append(item)
                completed_rows.add(item["row_number"])

    service = ComplianceWorkflowService(
        SqlAlchemyKnowledgeRepository(session),
        SqlAlchemyWorkflowRepository(session),
    )
    for row in rows:
        if row["row_number"] in completed_rows:
            continue
        run = service.execute(
            product={
                "external_id": row["external_id"],
                "title": row["title"],
                "description": row["description"],
                "category": row["category"],
                "attributes": {},
            },
            markets=row["markets"],
            category=row["category"],
            channel="all",
            as_of=None,
        )
        evidence_count = sum(
            len(item.get("candidate_evidence", []))
            for item in run.result_payload.get("markets", [])
        )
        result = {
            "row_number": row["row_number"],
            "external_id": row["external_id"],
            "workflow_id": run.id,
            "status": run.status.value,
            "evidence_count": evidence_count,
            "markets": ",".join(row["markets"]),
        }
        with checkpoint_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(result, ensure_ascii=False) + "\n")
        results.append(result)

    output_path = output_dir / "review-results.csv"
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "row_number", "external_id", "workflow_id", "status", "evidence_count", "markets"
            ],
        )
        writer.writeheader()
        writer.writerows(
            {key: _safe_cell(value) for key, value in item.items()} for item in results
        )
    return {
        "row_count": len(rows),
        "completed_count": len(results),
        "status_counts": {
            status: sum(item["status"] == status for item in results)
            for status in sorted({item["status"] for item in results})
        },
        "output_path": str(output_path.resolve()),
    }
